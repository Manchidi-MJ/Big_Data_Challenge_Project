"""Generate presentation/final_presentation.pptx from verified results.

Run from the project root:
    python presentation/build_pptx.py

Embeds the real figures in outputs/figures/ and the verified numbers. Speaker
notes include a short script for each of the three team members.
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "outputs", "figures")
OUT = os.path.join(ROOT, "presentation", "final_presentation.pptx")


def _title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    return slide


def _bullets_slide(prs, title, bullets, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = b
        p.font.size = Pt(18)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _figure_slide(prs, title, fig_name, bullets, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    fig_path = os.path.join(FIG, fig_name)
    if os.path.exists(fig_path):
        slide.shapes.add_picture(fig_path, Inches(0.4), Inches(1.5),
                                 height=Inches(4.6))
    # bullets on the right
    tx = slide.shapes.add_textbox(Inches(6.6), Inches(1.5), Inches(3.1), Inches(4.5))
    tf = tx.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(14)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def build() -> str:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 1 Title
    _title_slide(
        prs,
        "E-Commerce Customer Churn & Behaviour",
        "Big Data Analytics, RFM Segmentation & ML-Based Churn Prediction\n"
        "Mkhawana T. Hleketani  |  Burgos Carrera, M. de Los Angeles  |  Manchidi M. John",
    )

    # 2 Business context
    _bullets_slide(
        prs, "Business context & problem statement",
        [
            "E-commerce generates rich behavioural data on loyalty and churn.",
            "Retailers struggle to spot at-risk customers before they leave.",
            "Retention is an important lever (cost advantage is context-dependent).",
            "Goal: turn data into measurable findings and retention guidance.",
        ],
        "MA: Set up the motivation and state the problem — indicators exist but "
        "their link to churn isn't obvious without structured analysis.",
    )

    # 3 Research question & objectives
    _bullets_slide(
        prs, "Research question & objectives",
        [
            "Main Q: which behaviours relate to churn, and how well can ML predict it?",
            "SQ1: churn proportion; how churned vs retained differ.",
            "SQ2: which RFM segments are most at risk.",
            "SQ3: best model and most important features.",
            "Objectives: document data, clean (PySpark), EDA, RFM, model, recommend, limit.",
        ],
        "JM: Frame the questions; preview that the ML answer is nuanced due to a "
        "data issue we found.",
    )

    # 4 Dataset
    _bullets_slide(
        prs, "Dataset, provenance & limitations",
        [
            "Kaggle: E-Commerce Customer Churn 2025 - 10K (MIT licence).",
            "Verified 10,000 rows x 15 columns; one row per customer.",
            "No missing values, no duplicates. Churn rate 37.7%.",
            "Real schema differed from the brief's provisional list.",
            "Real vs synthetic: UNVERIFIED (author gives no description).",
        ],
        "TH: We verified the real schema and are explicit about not confirming "
        "provenance.",
    )

    # 5 Data engineering
    _bullets_slide(
        prs, "Data engineering & PySpark pipeline",
        [
            "Reproducible: read -> dedupe -> impute -> normalise -> derive tenure -> persist.",
            "Pandas and PySpark paths agree (10,000 -> 10,000 rows; nothing dropped).",
            "10k rows are NOT 'big data' - PySpark demonstrates a scalable DESIGN.",
            "Same code would run on a cluster for millions of customers.",
        ],
        "TH: Explain learn-once-run-anywhere design and the Windows winutils/Parquet note.",
    )

    # 6 EDA
    _figure_slide(
        prs, "EDA & churn distribution", "01_churn_distribution.png",
        [
            "37.7% churn (3,771 of 10,000).",
            "Churned vs retained differ almost only in recency: ~301 vs ~9 days.",
            "Spend, orders, age, device, category: nearly identical across groups.",
        ],
        "MA: Recency dominates; demographics and spend do not separate churners.",
    )

    # 7 RFM
    _figure_slide(
        prs, "RFM segmentation", "04_categorical_churn_rates.png",
        [
            "True RFM: Recency=days_since, Frequency=orders, Monetary=spend.",
            "'Needs Attention' (31.8%) and 'At Risk' (8.3%): ~94% churn.",
            "Recent segments (Champions, Loyal, ...): ~0% churn.",
        ],
        "MA: Explain 5-bin scoring with recency reversed; segments split by recency.",
    )

    # 8 ML approach
    _bullets_slide(
        prs, "Machine-learning approach",
        [
            "Binary classification of 'churned'.",
            "Preprocessing inside sklearn Pipelines (fit on train folds only).",
            "Models: DummyClassifier, Logistic Regression, Random Forest.",
            "Stratified 80/20 split, seed 42, 5-fold CV for selection.",
            "Two experiments: A = all features, B = leakage-resistant (primary).",
        ],
        "JM: Emphasise no leakage from scaling/encoding and why two experiments.",
    )

    # 9 Leakage
    _figure_slide(
        prs, "Leakage finding (key result)", "05_correlation_heatmap.png",
        [
            "churned == (days_since_last_purchase > 90) at 100% accuracy.",
            "This field DEFINES the label -> target-defining leakage.",
            "Dates encode the same recency -> also leakage.",
        ],
        "JM: This is the most important result; the 'best' naive model is not real "
        "prediction.",
    )

    # 10 Evaluation
    _figure_slide(
        prs, "Model evaluation & comparison", "roc_experiment_B.png",
        [
            "Experiment A ~100% (because of leakage) - a sanity check, not a result.",
            "Experiment B: RF ROC-AUC 0.54, F1 0.37; LR F1 0.46.",
            "Barely above majority baseline (Dummy accuracy 0.623).",
            "Low recall + precision -> unfit for targeting.",
        ],
        "JM: Contrast A vs B; discuss false-negative cost; no honest predictive power.",
    )

    # 11 Interpretation + recommendations
    _figure_slide(
        prs, "Interpretation & retention recommendations", "importance_rf_expb.png",
        [
            "Exp B importances all tiny; no feature stands out.",
            "R1 Re-engage before day 90. R2 Protect Champions/Loyal.",
            "R3 Prioritise high-value 'At Risk'. R4 Don't deploy a predictive model yet.",
            "Associations, not causation.",
        ],
        "MA+JM: Keep findings separate from proposed actions; no causal claims.",
    )

    # 12 Conclusion
    _bullets_slide(
        prs, "Conclusion, limitations & future work",
        [
            "Delivered a tested, reproducible, leakage-aware pipeline end-to-end.",
            "Headline: churn is a recency rule; apparent 100% models restate the label.",
            "Dataset describes the base but can't predict FUTURE churn.",
            "Future: event-level forward-labelled data, A/B test, scale-test Spark.",
            "Thank you - questions?",
        ],
        "All: Summarise the honest story and invite questions.",
    )

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote presentation -> {path}")
